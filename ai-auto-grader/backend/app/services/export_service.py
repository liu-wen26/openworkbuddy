"""导出服务（F10 系列）。

- F10-01 成绩明细表导出 Excel
- F10-02 答卷/错题图片导出（zip，可选水印）
- F10-03 学情报表导出（Excel 多 sheet / 可打印 HTML，HTML 可另存为 PDF）

所有导出以内存字节流返回，不落临时文件。
"""

import csv
import io
import logging
import zipfile
from typing import Dict, List, Optional
from uuid import UUID

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.answer_block import AnswerBlock
from app.models.exam import Exam
from app.models.student import ExamStudent, Student
from app.services import analytics_service, settings_service
from app.utils.file_storage import absolute_path

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)


def _exam_or_404(db: Session, exam_id: UUID) -> Exam:
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="考试不存在")
    return exam


def _autosize(ws) -> None:
    for column in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(length + 4, 10), 40)


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _workbook_bytes(wb: Workbook) -> bytes:
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


# ---------------- F10-01 成绩明细表 ----------------

def build_grade_detail(db: Session, exam_id: UUID) -> tuple[bytes, str]:
    exam = _exam_or_404(db, exam_id)

    roster = (
        db.query(ExamStudent, Student)
        .join(Student, ExamStudent.student_id == Student.id)
        .filter(ExamStudent.exam_id == exam_id)
        .all()
    )
    roster_map = {es.student_id: st for es, st in roster}
    totals = analytics_service._student_totals(analytics_service._collect(db, exam_id))
    qreport = analytics_service.question_report(db, exam_id)
    questions = qreport["questions"]

    ranked = sorted(totals.items(), key=lambda kv: kv[1]["total"], reverse=True)
    rank_grade = {sid: i + 1 for i, (sid, _) in enumerate(ranked)}

    # 班级内排名
    by_class: Dict[str, List[UUID]] = {}
    for sid, _ in ranked:
        cls = roster_map[sid].class_name or "未分班" if sid in roster_map else "未分班"
        by_class.setdefault(cls, []).append(sid)
    rank_class = {sid: i + 1 for ids in by_class.values() for i, sid in enumerate(ids)}

    # 逐题得分
    per_question: Dict[str, Dict[UUID, float]] = {q["question_number"]: {} for q in questions}
    for r in analytics_service._collect(db, exam_id)["choice"]:
        if r.student_id and r.question_number in per_question:
            per_question[r.question_number][r.student_id] = float(r.score or 0)
    for r in analytics_service._collect(db, exam_id)["subjective"]:
        score = analytics_service._effective_subjective(r)
        if r.student_id and r.question_number in per_question and score is not None:
            per_question[r.question_number][r.student_id] = score

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩明细"
    headers = ["考号", "姓名", "班级", "选择题得分", "非选择题得分", "总分", "年级排名", "班级排名"]
    headers += [f"第{q['question_number']}题({q['max_score']:g}分)" for q in questions]
    headers += ["是否缺考"]
    ws.append(headers)

    for es, st in sorted(roster, key=lambda x: (x[1].class_name or "", x[1].exam_number)):
        b = totals.get(st.id)
        row = [
            st.exam_number,
            st.name,
            st.class_name or "未分班",
            round(b["choice"], 2) if b else 0,
            round(b["subjective"], 2) if b else 0,
            b["total"] if b else 0,
            rank_grade.get(st.id),
            rank_class.get(st.id),
        ]
        row += [per_question.get(q["question_number"], {}).get(st.id) for q in questions]
        row.append("是" if es.is_absent else "否")
        ws.append(row)

    # 未在花名册但有成绩的考生
    extra = [sid for sid in totals if sid not in roster_map]
    for sid in extra:
        b = totals[sid]
        row = [f"(未匹配{sid})", "", "", round(b["choice"], 2), round(b["subjective"], 2), b["total"],
               rank_grade.get(sid), None]
        row += [per_question.get(q["question_number"], {}).get(sid) for q in questions]
        row.append("否")
        ws.append(row)

    _style_header(ws)
    _autosize(ws)
    ws.freeze_panes = "D2"

    filename = f"{exam.name}_成绩明细.xlsx"
    return _workbook_bytes(wb), filename


# ---------------- F10-03 学情报表 ----------------

def build_report_workbook(db: Session, exam_id: UUID) -> tuple[bytes, str]:
    exam = _exam_or_404(db, exam_id)
    overview = analytics_service.overview(db, exam_id)
    classes = analytics_service.class_report(db, exam_id)
    qreport = analytics_service.question_report(db, exam_id)
    knowledge = analytics_service.knowledge_report(db, exam_id)

    wb = Workbook()

    ws = wb.active
    ws.title = "年级总览"
    ws.append([f"{exam.name} 学情分析总览"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["指标", "数值"])
    for label, value in [
        ("科目", overview["subject"]), ("年级", overview["grade"]),
        ("满分", overview["total_score"]), ("及格线", overview["pass_score"]),
        ("优秀线", overview["excellent_score"]),
        ("应考人数", overview["roster_count"]), ("实考人数", overview["appeared_count"]),
        ("缺考人数", overview["absent_count"]), ("平均分", overview["avg_score"]),
        ("中位数", overview["median_score"]), ("最高分", overview["max_score"]),
        ("最低分", overview["min_score"]), ("标准差", overview["score_std"]),
        ("及格率", f"{overview['pass_rate'] * 100:.1f}%"),
        ("优秀率", f"{overview['excellent_rate'] * 100:.1f}%"),
    ]:
        ws.append([label, value])
    _style_header(ws)
    _autosize(ws)

    ws2 = wb.create_sheet("分数分布")
    ws2.append(["分数段", "人数"])
    for b in overview["distribution"]:
        ws2.append([b["label"], b["count"]])
    _style_header(ws2)
    _autosize(ws2)

    ws3 = wb.create_sheet("班级学情")
    ws3.append(["班级", "应考", "实考", "缺考", "平均分", "中位数", "最高分", "最低分",
                "及格率", "优秀率", "选择题得分率", "非选择题得分率"])
    for c in classes["classes"]:
        ws3.append([
            c["class_name"], c["student_count"], c["appeared"], c["absent"], c["avg_score"],
            c["median_score"], c["max_score"], c["min_score"],
            f"{c['pass_rate'] * 100:.1f}%", f"{c['excellent_rate'] * 100:.1f}%",
            f"{c.get('choice_rate', 0) * 100:.1f}%", f"{c.get('subjective_rate', 0) * 100:.1f}%",
        ])
    _style_header(ws3)
    _autosize(ws3)

    ws4 = wb.create_sheet("题目分析")
    ws4.append(["题号", "题型", "满分", "已评分", "平均分", "得分率", "难度", "区分度", "正确率", "知识点"])
    for q in qreport["questions"]:
        ws4.append([
            q["question_number"],
            "选择题" if q["question_type"] == "choice" else "非选择题",
            q["max_score"], q["graded"], q["avg_score"],
            f"{q['score_rate'] * 100:.1f}%", q["difficulty"], q["discrimination"],
            f"{q['correct_rate'] * 100:.1f}%" if q["correct_rate"] is not None else "—",
            "、".join(q["knowledge_tags"]),
        ])
    _style_header(ws4)
    _autosize(ws4)

    ws5 = wb.create_sheet("知识点掌握")
    ws5.append(["知识点", "关联题目", "得分率", "掌握等级"])
    level_label = {"mastered": "已掌握", "basic": "基本掌握", "weak": "薄弱"}
    for t in knowledge["tags"]:
        ws5.append([
            t["knowledge_tag"], "、".join(t["question_numbers"]),
            f"{t['avg_score_rate'] * 100:.1f}%", level_label.get(t["mastery_level"], t["mastery_level"]),
        ])
    _style_header(ws5)
    _autosize(ws5)

    filename = f"{exam.name}_学情分析报表.xlsx"
    return _workbook_bytes(wb), filename


def build_report_html(db: Session, exam_id: UUID) -> tuple[bytes, str]:
    """可打印的 HTML 学情报告（浏览器可「打印 → 另存为 PDF」）。"""
    exam = _exam_or_404(db, exam_id)
    overview = analytics_service.overview(db, exam_id)
    classes = analytics_service.class_report(db, exam_id)["classes"]
    qreport = analytics_service.question_report(db, exam_id)
    knowledge = analytics_service.knowledge_report(db, exam_id)["tags"]
    watermark = settings_service.get_watermark(db)

    def rows(items, keys, headers):
        head = "".join(f"<th>{h}</th>" for h in headers)
        body = ""
        for item in items:
            body += "<tr>" + "".join(f"<td>{item.get(k, '')}</td>" for k in keys) + "</tr>"
        return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"

    level_label = {"mastered": "已掌握", "basic": "基本掌握", "weak": "薄弱"}
    class_rows = [
        {**c, "pass_rate": f"{c['pass_rate'] * 100:.1f}%", "excellent_rate": f"{c['excellent_rate'] * 100:.1f}%"}
        for c in classes
    ]
    question_rows = [
        {**q, "score_rate": f"{q['score_rate'] * 100:.1f}%",
         "correct_rate": f"{q['correct_rate'] * 100:.1f}%" if q["correct_rate"] is not None else "—",
         "knowledge_tags": "、".join(q["knowledge_tags"])}
        for q in qreport["questions"]
    ]
    knowledge_rows = [
        {**t, "avg_score_rate": f"{t['avg_score_rate'] * 100:.1f}%",
         "question_numbers": "、".join(t["question_numbers"]),
         "mastery_level": level_label.get(t["mastery_level"], t["mastery_level"])}
        for t in knowledge
    ]

    wm_css = ""
    wm_html = ""
    if watermark.get("enabled"):
        wm_html = f'<div class="watermark">{watermark.get("text", "")}</div>'
        wm_css = (
            f".watermark{{position:fixed;top:40%;left:5%;right:5%;text-align:center;"
            f"font-size:{watermark.get('font_size', 16) * 3}px;color:{watermark.get('color', '#909399')};"
            f"opacity:{watermark.get('opacity', 0.12)};transform:rotate({watermark.get('rotate', -25)}deg);"
            f"pointer-events:none;z-index:0;}}"
        )

    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8" />
<title>{exam.name} 学情分析报告</title>
<style>
body{{font-family:"Microsoft YaHei",sans-serif;margin:32px;color:#303133;position:relative;}}
h1{{font-size:22px;}} h2{{font-size:17px;margin-top:28px;border-left:4px solid #409eff;padding-left:8px;}}
table{{border-collapse:collapse;width:100%;margin-top:10px;font-size:13px;position:relative;z-index:1;}}
th,td{{border:1px solid #dcdfe6;padding:6px 8px;text-align:center;}}
th{{background:#eef2f9;}}
.cards{{display:flex;flex-wrap:wrap;gap:12px;margin-top:12px;}}
.card{{border:1px solid #ebeef5;border-radius:6px;padding:10px 16px;min-width:120px;text-align:center;}}
.card .v{{font-size:20px;font-weight:bold;}} .card .l{{color:#909399;font-size:12px;margin-top:4px;}}
{wm_css}
</style></head><body>
{wm_html}
<h1>{exam.name} 学情分析报告</h1>
<p>{overview['subject']} · {overview['grade']} · 满分 {overview['total_score']:g} · 及格线 {overview['pass_score']:g} · 优秀线 {overview['excellent_score']:g}</p>
<div class="cards">
  <div class="card"><div class="v">{overview['appeared_count']}/{overview['roster_count']}</div><div class="l">实考/应考</div></div>
  <div class="card"><div class="v">{overview['avg_score']}</div><div class="l">平均分</div></div>
  <div class="card"><div class="v">{overview['max_score']}</div><div class="l">最高分</div></div>
  <div class="card"><div class="v">{overview['min_score']}</div><div class="l">最低分</div></div>
  <div class="card"><div class="v">{overview['pass_rate']*100:.1f}%</div><div class="l">及格率</div></div>
  <div class="card"><div class="v">{overview['excellent_rate']*100:.1f}%</div><div class="l">优秀率</div></div>
</div>
<h2>分数分布</h2>
{rows(overview['distribution'], ['label', 'count'], ['分数段', '人数'])}
<h2>班级学情</h2>
{rows(class_rows, ['class_name', 'student_count', 'appeared', 'absent', 'avg_score', 'median_score', 'pass_rate', 'excellent_rate'], ['班级', '应考', '实考', '缺考', '平均分', '中位数', '及格率', '优秀率'])}
<h2>题目分析</h2>
{rows(question_rows, ['question_number', 'max_score', 'graded', 'avg_score', 'score_rate', 'difficulty', 'discrimination', 'correct_rate', 'knowledge_tags'], ['题号', '满分', '已评分', '平均分', '得分率', '难度', '区分度', '正确率', '知识点'])}
<h2>知识点掌握</h2>
{rows(knowledge_rows, ['knowledge_tag', 'question_numbers', 'avg_score_rate', 'mastery_level'], ['知识点', '关联题目', '得分率', '掌握等级'])}
</body></html>"""
    return html.encode("utf-8"), f"{exam.name}_学情分析报告.html"


# ---------------- F10-02 答卷/错题图片导出 ----------------

def _apply_watermark(image_bytes: bytes, watermark: dict) -> bytes:
    if not watermark.get("enabled"):
        return image_bytes
    try:
        from PIL import Image, ImageDraw, ImageFont

        img = Image.open(io.BytesIO(image_bytes)).convert("RGBA")
        overlay = Image.new("RGBA", img.size, (255, 255, 255, 0))
        draw = ImageDraw.Draw(overlay)
        text = watermark.get("text") or "AI自动阅卷系统"
        font_size = max(12, int(watermark.get("font_size", 16) * max(img.size) / 400))
        try:
            font = ImageFont.truetype("DejaVuSans.ttf", font_size)
        except OSError:
            font = ImageFont.load_default()
        color = watermark.get("color", "#909399").lstrip("#")
        try:
            rgb = tuple(int(color[i:i + 2], 16) for i in (0, 2, 4))
        except (ValueError, IndexError):
            rgb = (144, 147, 153)
        alpha = int(max(0.0, min(1.0, watermark.get("opacity", 0.12))) * 255)
        pos = (int(img.size[0] * 0.05), int(img.size[1] * 0.4))
        draw.text(pos, text, font=font, fill=rgb + (alpha,))
        merged = Image.alpha_composite(img, overlay).convert("RGB")
        out = io.BytesIO()
        merged.save(out, format="PNG")
        return out.getvalue()
    except Exception as exc:  # noqa: BLE001  水印失败不影响导出
        logger.warning("水印绘制失败: %s", exc)
        return image_bytes


def build_answer_images_zip(
    db: Session,
    exam_id: UUID,
    *,
    student_id: Optional[UUID] = None,
    question_number: Optional[str] = None,
    mark: Optional[str] = None,
):
    exam = _exam_or_404(db, exam_id)
    watermark = settings_service.get_watermark(db)

    query = db.query(AnswerBlock).filter(AnswerBlock.exam_id == exam_id)
    if student_id:
        query = query.filter(AnswerBlock.student_id == student_id)
    if question_number:
        query = query.filter(AnswerBlock.question_number == question_number)
    blocks = query.all()

    # mark 过滤基于非选择题标记
    if mark:
        marked_ids = {
            r.block_id
            for r in analytics_service._collect(db, exam_id)["subjective"]
            if r.mark == mark
        }
        blocks = [b for b in blocks if b.id in marked_ids]

    student_map = {s.id: s for s in db.query(Student).all()}

    stream = io.BytesIO()
    manifest_rows: List[list] = [["文件名", "考号", "姓名", "班级", "题号", "类型", "水印"]]
    count = 0
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as zf:
        for block in blocks:
            if not block.image_path:
                continue
            path = absolute_path(block.image_path)
            if not path.exists():
                continue
            data = path.read_bytes()
            data = _apply_watermark(data, watermark)
            st = student_map.get(block.student_id)
            name = f"{st.exam_number if st else 'unknown'}_{st.name if st else '未匹配'}_第{block.question_number or '-'}题_{block.id}.png"
            zf.writestr(name, data)
            manifest_rows.append([
                name,
                st.exam_number if st else "",
                st.name if st else "",
                st.class_name if st else "",
                block.question_number or "",
                block.block_type,
                "是" if watermark.get("enabled") else "否",
            ])
            count += 1

        buf = io.StringIO()
        csv.writer(buf).writerows(manifest_rows)
        zf.writestr("manifest.csv", "\ufeff" + buf.getvalue())

    if count == 0:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="没有符合条件的答卷图片")

    scope = "错题" if mark == "typical_error" else "答卷"
    return stream.getvalue(), f"{exam.name}_{scope}图片_{count}张.zip"