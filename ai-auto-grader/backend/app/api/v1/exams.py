import io
from typing import List, Optional
from uuid import UUID
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.db.base import get_db
from app.core.permissions import require_permission
from app.models.user import User
from app.models.exam import Exam, ExamTeacher
from app.models.student import Student, ExamStudent
from app.models.template import AnswerCardTemplate
from app.models.template_config import TemplateRegion
from app.schemas.exam import (
    ExamCreate,
    ExamUpdate,
    ExamOut,
    ExamTeachersUpdate,
    ExamOriginalPaperOut,
)
from app.schemas.student import StudentImportResult, ExamStudentOut
from app.schemas.user import UserOut
from app.services import analytics_service
from app.utils.file_storage import save_original_paper, delete_original_paper
from app.core.config import get_settings

router = APIRouter(prefix="/exams", tags=["Exams"])


def datetime_now():
    from datetime import datetime, timezone
    return datetime.now(timezone.utc)


def check_exam_exists(db: Session, exam_id: UUID) -> Exam:
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Exam not found")
    return exam


def check_exam_modifiable(exam: Exam) -> None:
    if exam.status in ("locked", "archived"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Exam is locked or archived")


@router.post("", response_model=ExamOut)
def create_exam(
    payload: ExamCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:create")),
):
    if payload.answer_card_template_id:
        tpl = db.query(AnswerCardTemplate).filter(AnswerCardTemplate.id == payload.answer_card_template_id).first()
        if not tpl:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Template not found")
    exam = Exam(
        name=payload.name,
        subject=payload.subject,
        grade=payload.grade,
        exam_type=payload.exam_type,
        total_score=payload.total_score,
        pass_score=payload.pass_score,
        excellent_score=payload.excellent_score,
        answer_card_template_id=payload.answer_card_template_id,
        grading_start_at=payload.grading_start_at,
        grading_end_at=payload.grading_end_at,
        created_by=current_user.id,
        status="draft",
    )
    db.add(exam)
    db.commit()
    db.refresh(exam)
    return exam


@router.get("", response_model=List[ExamOut])
def list_exams(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:list")),
):
    return db.query(Exam).order_by(Exam.created_at.desc()).all()


@router.get("/{exam_id}", response_model=ExamOut)
def get_exam(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:view")),
):
    return check_exam_exists(db, exam_id)


@router.put("/{exam_id}", response_model=ExamOut)
def update_exam(
    exam_id: UUID,
    payload: ExamUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:update")),
):
    exam = check_exam_exists(db, exam_id)
    check_exam_modifiable(exam)
    data = payload.model_dump(exclude_unset=True)
    for key, value in data.items():
        setattr(exam, key, value)
    db.commit()
    db.refresh(exam)
    return exam


@router.delete("/{exam_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_exam(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:delete")),
):
    exam = check_exam_exists(db, exam_id)
    check_exam_modifiable(exam)
    db.query(ExamStudent).filter(ExamStudent.exam_id == exam_id).delete()
    db.query(ExamTeacher).filter(ExamTeacher.exam_id == exam_id).delete()
    db.delete(exam)
    db.commit()
    return None


@router.post("/{exam_id}/original-paper", response_model=ExamOut)
def upload_original_paper(
    exam_id: UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("paper:upload")),
):
    exam = check_exam_exists(db, exam_id)
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only PDF files are allowed")
    relative_path = save_original_paper(exam_id, file)
    exam.original_paper_path = relative_path
    exam.original_paper_uploaded_by = current_user.id
    exam.original_paper_uploaded_at = datetime_now()
    db.commit()
    db.refresh(exam)
    return exam


@router.delete("/{exam_id}/original-paper", response_model=ExamOut)
def remove_original_paper(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("paper:delete")),
):
    exam = check_exam_exists(db, exam_id)
    check_exam_modifiable(exam)
    delete_original_paper(exam_id)
    exam.original_paper_path = None
    exam.original_paper_uploaded_by = None
    exam.original_paper_uploaded_at = None
    db.commit()
    db.refresh(exam)
    return exam


@router.get("/{exam_id}/original-paper/download")
def download_original_paper(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("paper:view")),
):
    from fastapi.responses import FileResponse
    from pathlib import Path

    exam = check_exam_exists(db, exam_id)
    if not exam.original_paper_path:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Original paper not uploaded")
    full_path = Path(get_settings().STORAGE_ROOT) / exam.original_paper_path
    if not full_path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="File not found on disk")
    return FileResponse(
        path=str(full_path),
        filename=f"original_paper_{exam_id}.pdf",
        media_type="application/pdf",
    )


@router.get("/{exam_id}/original-paper/preview")
def preview_original_paper(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("paper:view")),
):
    """原试卷在线预览（inline，供浏览器内嵌查看）。"""
    from fastapi.responses import FileResponse
    from pathlib import Path

    exam = check_exam_exists(db, exam_id)
    if not exam.original_paper_path:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Original paper not uploaded")
    full_path = Path(get_settings().STORAGE_ROOT) / exam.original_paper_path
    if not full_path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="File not found on disk")
    return FileResponse(path=str(full_path), media_type="application/pdf")


@router.get("/{exam_id}/original-paper", response_model=ExamOriginalPaperOut)
def get_original_paper_info(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("paper:view")),
):
    exam = check_exam_exists(db, exam_id)
    return ExamOriginalPaperOut(
        original_paper_path=exam.original_paper_path,
        original_paper_uploaded_by=exam.original_paper_uploaded_by,
        original_paper_uploaded_at=exam.original_paper_uploaded_at,
    )


@router.post("/{exam_id}/students/import", response_model=StudentImportResult)
def import_students(
    exam_id: UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("student:import")),
):
    exam = check_exam_exists(db, exam_id)
    check_exam_modifiable(exam)

    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only Excel files are allowed")

    contents = file.file.read()
    workbook = load_workbook(filename=io.BytesIO(contents))
    sheet = workbook.active

    errors: List[str] = []
    success = 0
    total = 0

    # Expected header: 考号, 姓名, 班级
    headers = [cell.value for cell in sheet[1]]
    expected = ["考号", "姓名", "班级"]
    if headers[:3] != expected:
        errors.append(f"表头不正确，应为：{expected}，实际：{headers[:3]}")
        return StudentImportResult(total=0, success=0, failed=0, errors=errors)

    for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        total += 1
        exam_number, name, class_name = row[0], row[1], row[2] if len(row) > 2 else None
        if not exam_number or not name:
            errors.append(f"第 {idx} 行：考号或姓名为空")
            continue
        exam_number = str(exam_number).strip()
        name = str(name).strip()
        class_name = str(class_name).strip() if class_name else None

        # Avoid duplicate within same exam
        existing = (
            db.query(ExamStudent)
            .join(Student)
            .filter(ExamStudent.exam_id == exam_id, Student.exam_number == exam_number)
            .first()
        )
        if existing:
            errors.append(f"第 {idx} 行：考号 {exam_number} 已存在")
            continue

        student = db.query(Student).filter(Student.exam_number == exam_number).first()
        if not student:
            student = Student(exam_number=exam_number, name=name, class_name=class_name)
            db.add(student)
            db.flush()
        else:
            student.name = name
            student.class_name = class_name

        exam_student = ExamStudent(exam_id=exam_id, student_id=student.id)
        db.add(exam_student)
        success += 1

    db.commit()
    return StudentImportResult(total=total, success=success, failed=total - success, errors=errors)


@router.get("/{exam_id}/students", response_model=List[ExamStudentOut])
def list_exam_students(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:view")),
):
    check_exam_exists(db, exam_id)
    rows = (
        db.query(ExamStudent, Student)
        .join(Student, ExamStudent.student_id == Student.id)
        .filter(ExamStudent.exam_id == exam_id)
        .all()
    )
    board = analytics_service.scoreboard(db, exam_id)
    result = []
    for es, s in rows:
        board_row = board.get(s.id, {})
        result.append(ExamStudentOut(
            id=es.id,
            exam_id=es.exam_id,
            student_id=s.id,
            exam_number=s.exam_number,
            name=s.name,
            class_name=s.class_name,
            is_absent=es.is_absent,
            total_score=board_row.get("total") if s.id in board else None,
            rank_in_grade=board_row.get("rank_in_grade"),
            rank_in_class=board_row.get("rank_in_class"),
        ))
    return result


@router.post("/{exam_id}/teachers", status_code=status.HTTP_204_NO_CONTENT)
def assign_teachers(
    exam_id: UUID,
    payload: ExamTeachersUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:update")),
):
    exam = check_exam_exists(db, exam_id)
    check_exam_modifiable(exam)
    db.query(ExamTeacher).filter(ExamTeacher.exam_id == exam_id).delete()
    for item in payload.teachers:
        teacher = db.query(User).filter(User.id == item.teacher_id).first()
        if not teacher:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Teacher {item.teacher_id} not found")
        db.add(ExamTeacher(exam_id=exam_id, teacher_id=item.teacher_id, role_in_exam=item.role_in_exam))
    db.commit()
    return None


@router.get("/{exam_id}/teachers", response_model=List[UserOut])
def list_exam_teachers(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:view")),
):
    check_exam_exists(db, exam_id)
    teacher_ids = [
        t.teacher_id for t in db.query(ExamTeacher).filter(ExamTeacher.exam_id == exam_id).all()
    ]
    if not teacher_ids:
        return []
    return db.query(User).filter(User.id.in_(teacher_ids)).all()


@router.post("/{exam_id}/copy", response_model=ExamOut)
def copy_exam(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:create")),
):
    source = check_exam_exists(db, exam_id)
    new_exam = Exam(
        name=f"{source.name} (复制)",
        subject=source.subject,
        grade=source.grade,
        exam_type=source.exam_type,
        total_score=source.total_score,
        pass_score=source.pass_score,
        excellent_score=source.excellent_score,
        answer_card_template_id=source.answer_card_template_id,
        grading_start_at=source.grading_start_at,
        grading_end_at=source.grading_end_at,
        created_by=current_user.id,
        status="draft",
    )
    db.add(new_exam)
    db.commit()
    db.refresh(new_exam)
    return new_exam


@router.post("/{exam_id}/lock", response_model=ExamOut)
def lock_exam(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:update")),
):
    exam = check_exam_exists(db, exam_id)
    if exam.status == "archived":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Archived exam cannot be locked")
    exam.status = "locked"
    db.commit()
    db.refresh(exam)
    return exam


@router.post("/{exam_id}/unlock", response_model=ExamOut)
def unlock_exam(
    exam_id: UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:update")),
):
    exam = check_exam_exists(db, exam_id)
    if exam.status != "locked":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Exam is not locked")
    exam.status = "grading"
    db.commit()
    db.refresh(exam)
    return exam


@router.get("/{exam_id}/validate-score")
def validate_total_score(
    exam_id: UUID,
    configured_total: Decimal,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("exam:view")),
):
    """F2-05 总分校验：比对考试配置总分与答题卡模板各题块分值合计。"""
    exam = check_exam_exists(db, exam_id)
    template_total: Optional[float] = None
    if exam.answer_card_template_id:
        rows = (
            db.query(TemplateRegion.max_score)
            .filter(
                TemplateRegion.template_id == exam.answer_card_template_id,
                TemplateRegion.region_type.in_(("choice", "subjective")),
            )
            .all()
        )
        if rows:
            template_total = round(sum(float(r[0] or 0) for r in rows), 2)

    configured = float(configured_total)
    valid = template_total is None or abs(configured - template_total) < 0.01
    return {
        "valid": valid,
        "configured_total": configured,
        "template_total": template_total,
        "difference": round(configured - template_total, 2) if template_total is not None else None,
    }
